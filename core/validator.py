"""Cross-sheet consistency validator for a generated (or reference) workbook.

The validator reads ONLY the .xlsx file, so it can audit any workbook that
follows the layout contract in :mod:`core.report_spec` - including the
original hand-built reference. Every check re-derives a value from one
sheet and compares it with what another sheet states.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from core import report_spec as spec
from core.normalizer import Normalizer

_GAP_RE = re.compile(r"^(?P<name>.+?) \((?P<pct>\d+)% comps\)$")
_COUNT_RE = re.compile(r"^(?P<n>\d+) of (?P<total>\d+) competitors$")
_RATE_RE = re.compile(r"^(?P<n>\d+) of (?P<total>\d+) \((?P<pct>\d+)%\)$")
_REP_RE = re.compile(
    r"^(?P<tier>CRITICAL|HEALTHY): (?P<rating>\d(?:\.\d)?) Stars \((?P<revs>\d+) revs\)"
)


@dataclass(frozen=True)
class Issue:
    """One validation finding."""

    level: str  # "FAIL" or "WARN"
    sheet: str
    cell: str
    message: str

    def __str__(self) -> str:
        return f"[{self.level}] {self.sheet}!{self.cell}: {self.message}"


@dataclass
class ValidationReport:
    """Outcome of :func:`validate_workbook`."""

    path: str
    issues: list[Issue] = field(default_factory=list)
    checks: int = 0

    @property
    def failures(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "FAIL"]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "WARN"]

    @property
    def ok(self) -> bool:
        return not self.failures


@dataclass
class _ProfileClient:
    row: int
    address: str
    rating: str
    reviews: str
    primary: str
    all_categories: list[str]
    programs: list[str]


class WorkbookValidator:
    """Stateful helper; use :func:`validate_workbook`."""

    def __init__(self, path: str | Path, threshold_pct: int = 50, rep_floor: float = 4.0) -> None:
        self.path = str(path)
        self.threshold = threshold_pct
        self.rep_floor = rep_floor
        self.report = ValidationReport(path=self.path)
        self.norm = Normalizer()
        self.wb = openpyxl.load_workbook(self.path)

    # ------------------------------------------------------------- utilities
    def _fail(self, sheet: str, cell: str, msg: str) -> None:
        self.report.issues.append(Issue("FAIL", sheet, cell, msg))

    def _warn(self, sheet: str, cell: str, msg: str) -> None:
        self.report.issues.append(Issue("WARN", sheet, cell, msg))

    def _check(self, cond: bool, sheet: str, cell: str, msg: str) -> bool:
        self.report.checks += 1
        if not cond:
            self._fail(sheet, cell, msg)
        return cond

    @staticmethod
    def _rows(ws: Worksheet) -> list[tuple[Cell, ...]]:
        out: list[tuple[Cell, ...]] = []
        for row in ws.iter_rows(min_row=spec.FIRST_DATA_ROW):
            if row and row[0].value not in (None, ""):
                out.append(row)
        return out

    @staticmethod
    def _text(cell: Cell) -> str:
        return "" if cell.value is None else str(cell.value).strip()

    @staticmethod
    def _fill_hex(cell: Cell) -> str | None:
        if cell.fill is None or cell.fill.fill_type != "solid":
            return None
        rgb = cell.fill.fgColor.rgb
        return str(rgb)[-6:].upper() if rgb else None

    @staticmethod
    def _split_list(text: str, sep: str) -> list[str]:
        if not text or text == spec.NONE_LABEL:
            return []
        return [t.strip() for t in text.split(sep) if t.strip()]

    # ---------------------------------------------------------------- driver
    def run(self) -> ValidationReport:
        """Execute every check and return the report."""
        if not self._check_structure():
            return self.report
        profiles = self._index_profiles()
        services = self._check_services(profiles)
        categories = self._check_categories(profiles)
        self._check_summary(profiles, services, categories)
        return self.report

    # -------------------------------------------------------------- structure
    def _check_structure(self) -> bool:
        names = tuple(self.wb.sheetnames)
        ok = self._check(
            names == spec.SHEET_ORDER, "workbook", "-", f"sheet order {names} != {spec.SHEET_ORDER}"
        )
        if not ok:
            return False
        for sheet, headers in spec.HEADERS.items():
            ws = self.wb[sheet]
            found = tuple(
                self._text(ws.cell(row=spec.HEADER_ROW, column=i + 1)) for i in range(len(headers))
            )
            extra = self._text(ws.cell(row=spec.HEADER_ROW, column=len(headers) + 1))
            self._check(found == headers, sheet, f"A{spec.HEADER_ROW}", f"headers differ: {found}")
            self._check(
                extra == "",
                sheet,
                f"{spec.HEADER_ROW}",
                f"unexpected extra header column: {extra!r}",
            )
            self._check(self._text(ws.cell(row=1, column=1)) != "", sheet, "A1", "missing title")
            self._check(self._text(ws.cell(row=2, column=1)) != "", sheet, "A2", "missing subtitle")
        return self.report.ok

    # --------------------------------------------------------------- profiles
    def _index_profiles(self) -> dict[str, _ProfileClient]:
        sheet = spec.SHEET_PROFILES
        ws = self.wb[sheet]
        clients: dict[str, _ProfileClient] = {}
        self.true_counts: dict[str, int] = {}
        self.sister_counts: dict[str, int] = {}
        for row in self._rows(ws):
            cls, school = self._text(row[0]), self._text(row[1])
            r = row[0].row
            if cls == spec.CLASSIFICATION_CLIENT:
                self._check(
                    school not in clients,
                    sheet,
                    f"A{r}",
                    f"duplicate CLIENT SCHOOL row for {school}",
                )
                cats = self._split_list(self._text(row[10]), spec.LIST_SEP)
                primary = self._text(row[9])
                self._check(
                    bool(cats) and cats[0] == primary,
                    sheet,
                    f"K{r}",
                    f"'All Categories' must start with primary category {primary!r}: {cats}",
                )
                self._check(
                    self._text(row[2]) == school,
                    sheet,
                    f"C{r}",
                    "client entity name != school name",
                )
                clients[school] = _ProfileClient(
                    row=r,
                    address=self._text(row[5]),
                    rating=self._text(row[7]),
                    reviews=self._text(row[8]),
                    primary=primary,
                    all_categories=cats,
                    programs=self._split_list(self._text(row[11]), spec.PROGRAM_SEP),
                )
            elif cls == spec.CLASSIFICATION_TRUE:
                self.true_counts[school] = self.true_counts.get(school, 0) + 1
                self._entity_checks(row)
            elif cls == spec.CLASSIFICATION_SISTER:
                self.sister_counts[school] = self.sister_counts.get(school, 0) + 1
                self._entity_checks(row)
            else:
                self._fail(sheet, f"A{r}", f"unknown classification {cls!r}")
        self._check(bool(clients), sheet, "A5", "no CLIENT SCHOOL rows found")
        return clients

    def _entity_checks(self, row: tuple[Cell, ...]) -> None:
        sheet = spec.SHEET_PROFILES
        r = row[0].row
        name = self._text(row[2])
        is_sister = self.norm.is_sister_brand(name, self._text(row[3]))
        cls = self._text(row[0])
        self._check(
            is_sister == (cls == spec.CLASSIFICATION_SISTER),
            sheet,
            f"A{r}",
            f"{name!r} classified {cls!r} but sister-brand filter says {is_sister}",
        )
        cats = self._split_list(self._text(row[10]), spec.LIST_SEP)
        self._check(
            bool(cats) and cats[0] == self._text(row[9]),
            sheet,
            f"K{r}",
            "'All Categories' must start with the primary category",
        )
        for col, label in ((6, "distance"), (7, "rating"), (8, "review count")):
            txt = self._text(row[col])
            try:
                float(txt)
            except ValueError:
                self._fail(sheet, f"{chr(65 + col)}{r}", f"{label} is not numeric: {txt!r}")
            else:
                self.report.checks += 1

    # --------------------------------------------------------------- services
    def _check_services(
        self, profiles: dict[str, _ProfileClient]
    ) -> dict[str, list[dict[str, object]]]:
        sheet = spec.SHEET_SERVICES
        ws = self.wb[sheet]
        per_school: dict[str, list[dict[str, object]]] = {}
        for row in self._rows(ws):
            r = row[0].row
            school = self._text(row[0])
            svc = self._text(row[1])
            web = self._text(row[2]) == "YES"
            gbp_txt = self._text(row[3])
            gbp = gbp_txt == "YES"
            self._check(
                gbp_txt in ("YES", "NO (MISSING ON GBP)"),
                sheet,
                f"D{r}",
                f"unexpected GBP label {gbp_txt!r}",
            )
            m = _COUNT_RE.match(self._text(row[5]))
            if not self._check(
                m is not None, sheet, f"F{r}", f"cannot parse count {self._text(row[5])!r}"
            ):
                continue
            assert m is not None
            n, total = int(m["n"]), int(m["total"])
            expected_total = self.true_counts.get(school, 0)
            self._check(
                total == expected_total,
                sheet,
                f"F{r}",
                f"denominator {total} != {expected_total} TRUE COMPETITOR rows "
                "in Raw Profile Details",
            )
            pct_txt = self._text(row[4])
            expected_pct = f"{self.norm.pct(n, total)}%"
            self._check(
                pct_txt == expected_pct,
                sheet,
                f"E{r}",
                f"{pct_txt} != {expected_pct} (={n}/{total} half-up)",
            )
            pct = int(pct_txt.rstrip("%")) if pct_txt.rstrip("%").isdigit() else -1
            tier = self._service_tier(web, gbp, n, pct)
            finding = self._text(row[6])
            self._check(
                finding.startswith(self._service_prefix(tier)),
                sheet,
                f"G{r}",
                f"finding {finding[:40]!r} inconsistent with web={web} gbp={gbp} "
                f"pct={pct} -> {tier}",
            )
            fill_hex = self._fill_hex(row[3])
            self._check(
                fill_hex == spec.SERVICE_TIER_FILL[tier],
                sheet,
                f"D{r}",
                f"fill {fill_hex} != {spec.SERVICE_TIER_FILL[tier]} expected for {tier}",
            )
            per_school.setdefault(school, []).append(
                {"service": svc, "web": web, "gbp": gbp, "pct": pct, "tier": tier}
            )
        for school in per_school:
            self._check(
                school in profiles, sheet, "A", f"school {school!r} has no CLIENT SCHOOL row"
            )
        return per_school

    def _service_tier(self, web: bool, gbp: bool, count: int, pct: int) -> str:
        if web and not gbp and pct >= self.threshold:
            return "CRITICAL"
        if web and not gbp:
            return "INTERNAL"
        if not web and not gbp and pct >= self.threshold:
            return "MARKET"
        if gbp and pct >= self.threshold:
            return "DEFEND"
        if gbp and count == 0:
            return "USP"
        return "MONITOR"

    @staticmethod
    def _service_prefix(tier: str) -> str:
        return {
            "CRITICAL": "CRITICAL SYNC GAP",
            "INTERNAL": "INTERNAL SYNC GAP",
            "MARKET": "MARKET GAP",
            "DEFEND": "DEFEND",
            "USP": "DIFFERENTIATOR",
            "MONITOR": "MONITOR",
        }[tier]

    # ------------------------------------------------------------- categories
    def _check_categories(
        self, profiles: dict[str, _ProfileClient]
    ) -> dict[str, list[dict[str, object]]]:
        sheet = spec.SHEET_CATEGORIES
        ws = self.wb[sheet]
        per_school: dict[str, list[dict[str, object]]] = {}
        for row in self._rows(ws):
            r = row[0].row
            school = self._text(row[0])
            prof = profiles.get(school)
            if not self._check(
                prof is not None, sheet, f"A{r}", f"school {school!r} has no CLIENT SCHOOL row"
            ):
                continue
            assert prof is not None
            primary = self._text(row[1])
            additional = self._split_list(self._text(row[2]), spec.LIST_SEP)
            self._check(
                primary == prof.primary,
                sheet,
                f"B{r}",
                f"{primary!r} != profile primary {prof.primary!r}",
            )
            self._check(
                additional == prof.all_categories[1:],
                sheet,
                f"C{r}",
                f"{additional} != profile additional {prof.all_categories[1:]}",
            )
            cat = self._text(row[3])
            m = _RATE_RE.match(self._text(row[4]))
            if not self._check(
                m is not None, sheet, f"E{r}", f"cannot parse rate {self._text(row[4])!r}"
            ):
                continue
            assert m is not None
            n, total, pct = int(m["n"]), int(m["total"]), int(m["pct"])
            expected_total = self.true_counts.get(school, 0)
            self._check(
                total == expected_total,
                sheet,
                f"E{r}",
                f"denominator {total} != {expected_total} true competitors",
            )
            self._check(
                pct == self.norm.pct(n, total),
                sheet,
                f"E{r}",
                f"{pct}% != {self.norm.pct(n, total)}% (={n}/{total} half-up)",
            )
            covered = self.norm.category_key(cat) in {
                self.norm.category_key(c) for c in [primary, *additional]
            }
            tier = "COVERED" if covered else ("CRITICAL" if pct >= self.threshold else "CONSIDER")
            rec = self._text(row[6])
            self._check(
                rec.startswith(
                    {"COVERED": "COVERED", "CRITICAL": "CRITICAL GAP", "CONSIDER": "CONSIDER"}[tier]
                ),
                sheet,
                f"G{r}",
                f"recommendation {rec[:30]!r} inconsistent with covered={covered} "
                f"pct={pct} -> {tier}",
            )
            fill_hex = self._fill_hex(row[6])
            self._check(
                fill_hex == spec.CATEGORY_TIER_FILL[tier],
                sheet,
                f"G{r}",
                f"fill {fill_hex} != {spec.CATEGORY_TIER_FILL[tier]} expected for {tier}",
            )
            self._check(
                self._text(row[5]) == spec.GOOGLE_CLASSIFICATION,
                sheet,
                f"F{r}",
                "classification label",
            )
            per_school.setdefault(school, []).append({"category": cat, "pct": pct, "tier": tier})
        return per_school

    # ---------------------------------------------------------------- summary
    def _check_summary(
        self,
        profiles: dict[str, _ProfileClient],
        services: dict[str, list[dict[str, object]]],
        categories: dict[str, list[dict[str, object]]],
    ) -> None:
        sheet = spec.SHEET_SUMMARY
        ws = self.wb[sheet]
        seen: set[str] = set()
        for row in self._rows(ws):
            r = row[0].row
            school = self._text(row[2])
            seen.add(school)
            prof = profiles.get(school)
            if not self._check(
                prof is not None, sheet, f"C{r}", f"school {school!r} has no CLIENT SCHOOL row"
            ):
                continue
            assert prof is not None
            svc_rows = services.get(school, [])
            cat_rows = categories.get(school, [])

            self._check(
                self._text(row[3]) == prof.address,
                sheet,
                f"D{r}",
                f"address != Raw Profile {prof.address!r}",
            )
            expected_e = f"{prof.rating}* ({prof.reviews})"
            self._check(
                self._text(row[4]) == expected_e,
                sheet,
                f"E{r}",
                f"{self._text(row[4])!r} != {expected_e!r}",
            )

            true_n = self.true_counts.get(school, 0)
            self._check(
                str(row[5].value) == str(true_n),
                sheet,
                f"F{r}",
                f"{row[5].value} != {true_n} TRUE COMPETITOR rows",
            )
            self._check(
                str(row[6].value) == str(len(prof.programs)),
                sheet,
                f"G{r}",
                f"{row[6].value} != {len(prof.programs)} programs listed in Raw Profile Details",
            )
            yes_count = sum(1 for s in svc_rows if s["gbp"])
            h_val = str(row[7].value)
            if h_val.isdigit():
                h_int = int(h_val)
                self._check(
                    h_int >= yes_count,
                    sheet,
                    f"H{r}",
                    f"{h_int} chips claimed but matrix shows {yes_count} YES rows",
                )
                if h_int > yes_count:
                    self._warn(
                        sheet,
                        f"H{r}",
                        f"{h_int} chips claimed; only {yes_count} map to matrix rows "
                        "(unmapped or unmatched chips)",
                    )
            else:
                self._fail(sheet, f"H{r}", f"not an integer: {h_val!r}")

            # Internal sync gaps == matrix rows web YES / gbp NO
            expected_internal = {s["service"] for s in svc_rows if s["web"] and not s["gbp"]}
            found_internal = set(self._split_list(self._text(row[8]), spec.LIST_SEP))
            self._check(
                found_internal == expected_internal,
                sheet,
                f"I{r}",
                f"internal gaps {sorted(found_internal)} != matrix (web YES, GBP NO) "
                f"{sorted(expected_internal)}",
            )

            # External gaps == matrix rows gbp NO & pct >= threshold, with matching pct
            expected_external = {
                (s["service"], s["pct"])
                for s in svc_rows
                if not s["gbp"] and int(s["pct"]) >= self.threshold
            }
            found_external = self._parse_gaps(sheet, f"J{r}", self._text(row[9]))
            self._check(
                found_external == expected_external,
                sheet,
                f"J{r}",
                f"external gaps {sorted(found_external)} != matrix (GBP NO, >= "
                f"{self.threshold}%) {sorted(expected_external)}",
            )

            # Category gaps == CRITICAL category rows, with matching pct
            expected_cats = {(c["category"], c["pct"]) for c in cat_rows if c["tier"] == "CRITICAL"}
            found_cats = self._parse_gaps(sheet, f"K{r}", self._text(row[10]))
            self._check(
                found_cats == expected_cats,
                sheet,
                f"K{r}",
                f"category gaps {sorted(found_cats)} != CRITICAL category rows "
                f"{sorted(expected_cats)}",
            )

            # Reputation label
            rep = self._text(row[11])
            m = _REP_RE.match(rep)
            if self._check(m is not None, sheet, f"L{r}", f"cannot parse reputation {rep!r}"):
                assert m is not None
                self._check(
                    m["rating"] == prof.rating and m["revs"] == prof.reviews,
                    sheet,
                    f"L{r}",
                    f"{rep!r} != profile {prof.rating}/{prof.reviews}",
                )
                expected_tier = "CRITICAL" if float(prof.rating) < self.rep_floor else "HEALTHY"
                self._check(
                    m["tier"] == expected_tier,
                    sheet,
                    f"L{r}",
                    f"tier {m['tier']} != {expected_tier} for rating {prof.rating}",
                )
                fill_hex = self._fill_hex(row[11])
                self._check(
                    (fill_hex == spec.SOFT_RED) == (expected_tier == "CRITICAL"),
                    sheet,
                    f"L{r}",
                    f"fill {fill_hex} inconsistent with {expected_tier}",
                )
            self._check(self._text(row[12]) != "", sheet, f"M{r}", "empty action plan")
            self._check(
                self._fill_hex(row[8]) == spec.SOFT_RED,
                sheet,
                f"I{r}",
                "internal-gap cell must be soft red",
            )
            self._check(
                self._fill_hex(row[9]) == spec.SOFT_RED,
                sheet,
                f"J{r}",
                "external-gap cell must be soft red",
            )

        for school in profiles:
            self._check(
                school in seen,
                sheet,
                "C",
                f"CLIENT SCHOOL {school!r} missing from Executive Summary",
            )

    def _parse_gaps(self, sheet: str, cell: str, text: str) -> set[tuple[str, int]]:
        out: set[tuple[str, int]] = set()
        for item in self._split_list(text, spec.LIST_SEP):
            m = _GAP_RE.match(item)
            if not self._check(
                m is not None,
                sheet,
                cell,
                f"cannot parse gap entry {item!r} (expected 'Name (NN% comps)')",
            ):
                continue
            assert m is not None
            out.add((m["name"], int(m["pct"])))
        return out


def validate_workbook(
    path: str | Path, threshold_pct: int = 50, rep_floor: float = 4.0
) -> ValidationReport:
    """Validate ``path`` and return a :class:`ValidationReport`."""
    return WorkbookValidator(path, threshold_pct=threshold_pct, rep_floor=rep_floor).run()


def format_report(report: ValidationReport) -> str:
    """Render a report as plain text."""
    lines = [
        f"Workbook: {report.path}",
        f"Checks: {report.checks}  Failures: {len(report.failures)}  "
        f"Warnings: {len(report.warnings)}",
        f"Result: {'PASS' if report.ok else 'FAIL'}",
    ]
    lines.extend(str(i) for i in report.issues)
    return "\n".join(lines)
