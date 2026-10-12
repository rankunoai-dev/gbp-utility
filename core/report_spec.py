"""Workbook layout contract: sheet names, headers and style palette.

Both the reporter (writer) and the validator (reader) import from here, so
the format can only drift in one place.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET_SUMMARY = "Executive Summary & Gaps"
SHEET_SERVICES = "Service Comparison Matrix"
SHEET_CATEGORIES = "Category Gap Matrix"
SHEET_PROFILES = "Raw Profile Details"
SHEET_ORDER: tuple[str, ...] = (SHEET_SUMMARY, SHEET_SERVICES, SHEET_CATEGORIES, SHEET_PROFILES)

HEADER_ROW = 4
FIRST_DATA_ROW = 5

HEADERS_SUMMARY: tuple[str, ...] = (
    "School ID",
    "Brand",
    "Audited School Name",
    "Verified Physical Address",
    "GBP Star Rating (Reviews)",
    "True Competitors Audited",
    "Website Programs Offered",
    "GBP Service Chips Claimed",
    "Internal Website-to-GBP Sync Gaps (Offered by School, Missing on GBP)",
    "External Competitor Gaps (>=50% Competitors Offer)",
    "Missing GBP Categories (Gaps)",
    "Local Reputation Alert",
    "Priority Action Plan",
)
HEADERS_SERVICES: tuple[str, ...] = (
    "Audited Client School",
    "Program / Service Name",
    "Offered on School Website?",
    "Claimed as GBP Service Chip?",
    "True Competitor Adoption Rate",
    "Verified Competitors Offering",
    "Strategic Audit Finding",
)
HEADERS_CATEGORIES: tuple[str, ...] = (
    "Audited Client School",
    "Client Primary Category",
    "Client Additional Categories",
    "Competitor Category Discovered",
    "Competitor Adoption Rate (% & Count)",
    "Google Classification Type",
    "Actionable Recommendation",
)
HEADERS_PROFILES: tuple[str, ...] = (
    "Classification",
    "Audited School Location",
    "Brand / Entity Name",
    "Official Website URL",
    "Live Google Maps Profile URL",
    "Verified Street Address",
    "Distance (mi)",
    "Google Star Rating",
    "Google Review Count",
    "Primary Category",
    "All Categories",
    "Extracted Services & Programs",
)
HEADERS: dict[str, tuple[str, ...]] = {
    SHEET_SUMMARY: HEADERS_SUMMARY,
    SHEET_SERVICES: HEADERS_SERVICES,
    SHEET_CATEGORIES: HEADERS_CATEGORIES,
    SHEET_PROFILES: HEADERS_PROFILES,
}

CLASSIFICATION_CLIENT = "CLIENT SCHOOL"
CLASSIFICATION_TRUE = "TRUE COMPETITOR"
CLASSIFICATION_SISTER = "SISTER BRAND (FILTERED)"
GOOGLE_CLASSIFICATION = "Official Google Business Category"
NONE_LABEL = "None"
PROGRAM_SEP = " | "
LIST_SEP = ", "

# Palette ------------------------------------------------------------------
NAVY = "1F4E78"
DARK_TEAL = "005960"
CLIENT_BLUE = "D9E1F2"
SISTER_GREY = "EDEDED"
SOFT_RED = "FCE4D6"
SOFT_GREEN = "E2EFDA"
SOFT_YELLOW = "FFF2CC"
ALERT_RED = "C00000"
OK_GREEN = "375623"
MUTED_GREY = "595959"


def fill(hex_color: str) -> PatternFill:
    """Solid fill for ``hex_color``."""
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


class Styles:
    """Style objects shared by every sheet."""

    navy_fill = fill(NAVY)
    dark_teal_fill = fill(DARK_TEAL)
    client_blue_fill = fill(CLIENT_BLUE)
    sister_fill = fill(SISTER_GREY)
    soft_red_fill = fill(SOFT_RED)
    soft_green_fill = fill(SOFT_GREEN)
    soft_yellow_fill = fill(SOFT_YELLOW)

    font_title = Font(name="Calibri", size=15, bold=True, color=NAVY)
    font_subtitle = Font(name="Calibri", size=10, italic=True, color=MUTED_GREY)
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_muted = Font(name="Calibri", size=10, italic=True, color=MUTED_GREY)
    font_alert = Font(name="Calibri", size=10, bold=True, color=ALERT_RED)
    font_ok = Font(name="Calibri", size=10, bold=True, color=OK_GREEN)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)


# Fill colour expected on the "Claimed as GBP Service Chip?" cell per tier.
SERVICE_TIER_FILL: dict[str, str] = {
    "CRITICAL": SOFT_RED,
    "MARKET": SOFT_RED,
    "INTERNAL": SOFT_YELLOW,
    "DEFEND": SOFT_GREEN,
    "USP": SOFT_GREEN,
    "MONITOR": CLIENT_BLUE,
}
# Fill colour expected on the "Actionable Recommendation" cell per tier.
CATEGORY_TIER_FILL: dict[str, str] = {
    "CRITICAL": SOFT_RED,
    "CONSIDER": SOFT_YELLOW,
    "COVERED": SOFT_GREEN,
}
