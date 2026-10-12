import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
# Tab 1: Executive Summary & Gaps
ws1 = wb.active
ws1.title = "Executive Gap Analysis"

# Tab 2: Service Comparison Matrix
ws2 = wb.create_sheet(title="Service Comparison Matrix")

# Tab 3: Category Gap Matrix
ws3 = wb.create_sheet(title="Category Gap Matrix")

# Tab 4: Raw Competitor Data
ws4 = wb.create_sheet(title="Raw Profile Details")

print("Created 4 sheets successfully")
wb.save("test_out.xlsx")
