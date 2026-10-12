from pathlib import Path

import openpyxl
import pytest

from core import report_spec as spec
from core.gap_engine import GapEngine
from core.models import AuditDataset
from core.reporter import ReportGenerator
from core.validator import validate_workbook
from tests.conftest import REFERENCE_XLSX


@pytest.fixture(scope="module")
def generated(tmp_path_factory: pytest.TempPathFactory, pilot_dataset: AuditDataset) -> Path:
    engine = GapEngine(brand_short=pilot_dataset.client_brand_short)
    reporter = ReportGenerator(engine, brand_short=pilot_dataset.client_brand_short)
    analyses = [engine.analyze(s) for s in pilot_dataset.schools]
    out = tmp_path_factory.mktemp("wb") / "pilot.xlsx"
    reporter.generate_workbook(analyses, str(out))
    return out


def test_generated_workbook_passes_consistency_validation(generated: Path) -> None:
    report = validate_workbook(generated)
    assert report.failures == [], "\n".join(str(i) for i in report.failures)
    assert report.warnings == [], "\n".join(str(i) for i in report.warnings)
    assert report.checks > 300


def test_sheet_names_and_headers_match_contract(generated: Path) -> None:
    wb = openpyxl.load_workbook(generated)
    assert tuple(wb.sheetnames) == spec.SHEET_ORDER
    for sheet, headers in spec.HEADERS.items():
        ws = wb[sheet]
        found = tuple(ws.cell(row=spec.HEADER_ROW, column=i + 1).value for i in range(len(headers)))
        assert found == headers
        assert ws.cell(row=spec.HEADER_ROW, column=len(headers) + 1).value is None


@pytest.mark.skipif(not REFERENCE_XLSX.exists(), reason="reference workbook not present")
def test_layout_matches_reference_workbook(generated: Path) -> None:
    ref = openpyxl.load_workbook(REFERENCE_XLSX)
    gen = openpyxl.load_workbook(generated)
    assert gen.sheetnames == ref.sheetnames
    for name in ref.sheetnames:
        r, g = ref[name], gen[name]
        assert g.max_column == r.max_column, name
        for col in range(1, r.max_column + 1):
            rc, gc = r.cell(row=spec.HEADER_ROW, column=col), g.cell(row=spec.HEADER_ROW, column=col)
            assert gc.value == rc.value, f"{name} header col {col}"
            assert gc.fill.fgColor.rgb == rc.fill.fgColor.rgb, f"{name} header fill col {col}"
            assert gc.font.bold == rc.font.bold and gc.font.color.rgb == rc.font.color.rgb
        # Title / subtitle styling
        for row in (1, 2):
            assert g.cell(row=row, column=1).font.size == r.cell(row=row, column=1).font.size
            assert g.cell(row=row, column=1).font.color.rgb == r.cell(row=row, column=1).font.color.rgb
        # Data rows start at the same row and carry the same border/alignment rule
        assert g.cell(row=spec.FIRST_DATA_ROW, column=1).border.left.style == "thin"
        assert r.cell(row=spec.FIRST_DATA_ROW, column=1).border.left.style == "thin"


def test_executive_summary_values(generated: Path) -> None:
    ws = openpyxl.load_workbook(generated)[spec.SHEET_SUMMARY]
    fairport = [c.value for c in ws[5]]
    henrietta = [c.value for c in ws[6]]
    assert fairport[:8] == [
        "0002",
        "Childtime Learning Centers",
        "Childtime of Fairport",
        "585 Garnsey Rd, Fairport, NY 14450",
        "4.7* (12)",
        6,
        14,
        5,
    ]
    assert fairport[9] == "Preschool Program (100% comps), Summer Camp (83% comps)"
    assert fairport[10] == "None"
    assert fairport[11] == "HEALTHY: 4.7 Stars (12 revs)"
    assert henrietta[5:8] == [5, 14, 6]
    assert henrietta[10] == "Child care agency (80% comps)"
    assert henrietta[11].startswith("CRITICAL: 3.7 Stars (16 revs)")


def test_raw_profile_programs_are_in_taxonomy_order(generated: Path) -> None:
    ws = openpyxl.load_workbook(generated)[spec.SHEET_PROFILES]
    client_row = next(r for r in ws.iter_rows(min_row=5) if r[0].value == spec.CLASSIFICATION_CLIENT)
    assert client_row[11].value == (
        "Infants | Toddlers | Twos | Early Preschool | Preschool | Pre-K | Junior Kindergarten | "
        "Private Kindergarten | Before & After School | School Break Camps | Summer Camp | "
        "Drop-In Care | Grow Fit | Electives"
    )
    classifications = [r[0].value for r in ws.iter_rows(min_row=5) if r[0].value]
    assert classifications.count(spec.CLASSIFICATION_CLIENT) == 2
    assert classifications.count(spec.CLASSIFICATION_SISTER) == 2
    assert classifications.count(spec.CLASSIFICATION_TRUE) == 11


@pytest.mark.skipif(not REFERENCE_XLSX.exists(), reason="reference workbook not present")
def test_reference_workbook_is_known_to_be_inconsistent() -> None:
    """Documents why the engine must not simply copy the reference's numbers."""
    report = validate_workbook(REFERENCE_XLSX)
    assert not report.ok
    messages = " ".join(str(i) for i in report.failures)
    assert "66% != 67%" in messages
    assert "internal gaps" in messages
