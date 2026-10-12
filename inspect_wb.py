import openpyxl

wb = openpyxl.load_workbook('Childtime_GBP_Service_and_Category_Gap_Report.xlsx')
print('Sheet names:', wb.sheetnames)
ws = wb['Executive Summary & Gaps']
for row in ws.iter_rows(values_only=True):
    print(row[:5])
